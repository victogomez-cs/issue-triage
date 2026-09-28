"""Stage 4: score every issue and write the Excel workbook (plus a CSV of the same rows)."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from .profile import Profile
from .scoring import BUG_KINDS, KIND_DISPLAY, URGENCY_TIERS, build_context, parse_ts, score_issue

MAX_DUP_ROWS = 5000
MAX_TEAMS_IN_SUMMARY = 30

COLUMNS = [  # (header, key, width)
    ("#", "number", 9), ("Title", "title", 60), ("Kind", "kind", 18), ("Severity", "severity", 11),
    ("Urgency", "urgency", 10), ("Score", "score", 8), ("Priority", "priority", 9),
    ("Team / area", "team", 14), ("Triage status", "triage_status", 21), ("👍", "thumbs_up", 7),
    ("Comments", "comments", 10), ("Flags", "flags", 40), ("Possible duplicate of", "dup_of", 12),
    ("Dup similarity", "dup_sim", 11), ("Label-based kind", "label_kind", 15),
    ("AI kind", "ai_kind", 16), ("AI severity", "ai_severity", 11), ("AI confidence", "ai_confidence", 11),
    ("AI summary", "ai_summary", 50), ("Has repro", "has_repro", 11),
    ("Found in release", "release", 10), ("Created", "created", 12), ("Updated", "updated", 12),
    ("Days since update", "days_since_update", 10), ("Assignees", "assignees", 16),
    ("Score breakdown", "score_breakdown", 50), ("Labels", "labels", 60),
]
COL_INDEX = {key: i for i, (_, key, _) in enumerate(COLUMNS)}

FLAG_SUMMARY = [  # (label on the Summary sheet, wildcard pattern)
    ("Needs triage", "Needs triage"), ("Needs primary triage", "Needs primary triage"),
    ("Needs team triage", "Needs team triage"), ("Triaged, no priority", "Triaged, no priority"),
    ("Crash at P3", "Crash at P3"), ("Regression at P3", "Regression at P3"),
    ("Popular at P3", "thumbs-up) at P3"), ("AI: looks like a bug, labeled feature", "AI: looks like a bug"),
    ("AI: may not be a bug", "AI: may not be a bug"), ("AI: critical but P2/P3", "AI: critical but"),
    ("Possible duplicate", "Possible duplicate"), ("Stale 2+ years", "Stale 2+ years"),
]


def build_rows(blob: dict, profile: Profile, ai: dict, dups: dict):
    issues = blob["issues"]
    ctx = build_context(issues, profile, blob.get("exported_at"))
    dup_of = {}
    for pr in dups.get("pairs", []):
        newer, older, sim = max(pr["a"], pr["b"]), min(pr["a"], pr["b"]), pr["similarity"]
        if newer not in dup_of or dup_of[newer][1] < sim:
            dup_of[newer] = (older, sim)
    rows = []
    for i in issues:
        r = score_issue(i, ai.get(i["number"]), ctx, dup_of.get(i["number"]))
        r.update({
            "number": i["number"], "title": i["title"], "url": i["url"],
            "thumbs_up": i.get("thumbs_up", 0), "comments": i.get("comments", 0),
            "created": parse_ts(i["created_at"]).date(), "updated": parse_ts(i["updated_at"]).date(),
            "assignees": ", ".join(i.get("assignees") or []), "labels": ", ".join(i["labels"]),
        })
        rows.append(r)
    return rows, ctx


def write_report(blob: dict, profile: Profile, ai: dict, dups: dict, out_path: Path) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    from openpyxl.worksheet.table import Table, TableStyleInfo

    rows, ctx = build_rows(blob, profile, ai, dups)
    repo = blob.get("repo") or "repository"
    host = blob.get("host") or "github.com"
    titles = {i["number"]: i["title"] for i in blob["issues"]}
    w = profile.weights

    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True)
    link = Font(name="Arial", size=10, color="0563C1", underline="single")
    head_fill = PatternFill("solid", fgColor="1F3864")
    urgency_fills = {k: PatternFill("solid", fgColor=v) for k, v in
                     {"Critical": "F8CBAD", "High": "FCE4D6", "Medium": "FFF2CC", "Low": "E2EFDA"}.items()}

    wb = Workbook()

    def header(ws, names):
        ws.append(names)
        for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=len(names))):
            c.font, c.fill = head, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

    def write_table(ws, data, table_name, rank=False):
        names = (["Rank"] if rank else []) + [h for h, _, _ in COLUMNS]
        ncols, offset = len(names), 1 if rank else 0
        urg_col = offset + COL_INDEX["urgency"]
        date_cols = {offset + COL_INDEX[k] for k in ("created", "updated")}
        header(ws, names)
        for n, r in enumerate(data, 1):
            ws.append(([n] if rank else []) + [r.get(k) for _, k, _ in COLUMNS])
        for r, cells in zip(data, ws.iter_rows(min_row=2, max_row=len(data) + 1, max_col=ncols)):
            for idx, c in enumerate(cells):
                c.font = base
                if idx in date_cols:
                    c.number_format = "yyyy-mm-dd"
            cells[offset].hyperlink = r["url"]
            cells[offset].font = link
            if r.get("urgency") in urgency_fills:
                cells[urg_col].fill = urgency_fills[r["urgency"]]
        for col, (_, _, width) in enumerate(COLUMNS, offset + 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        if rank:
            ws.column_dimensions["A"].width = 7
        ws.freeze_panes = ws.cell(row=2, column=offset + 3)
        if data:
            t = Table(displayName=table_name, ref=f"A1:{get_column_letter(ncols)}{len(data) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
            ws.add_table(t)

    rank_key = lambda r: (-r["score"], -r["thumbs_up"], r["number"])  # noqa: E731
    ws_sum = wb.active
    ws_sum.title = "Summary"

    bugs = sorted([r for r in rows if r["kind"] in BUG_KINDS], key=rank_key)
    write_table(wb.create_sheet("Bugs Ranked"), bugs, "BugsRanked", rank=True)
    review = sorted([r for r in rows if r["flags"] and r["flags"] != "Stale 2+ years"], key=rank_key)
    write_table(wb.create_sheet("Review Queue"), review, "ReviewQueue")
    all_rows = sorted(rows, key=lambda r: r["number"])
    write_table(wb.create_sheet("All Issues"), all_rows, "AllIssues")

    # Duplicates
    ws_dup = wb.create_sheet("Duplicates")
    header(ws_dup, ["Newer issue", "Newer title", "Older issue", "Older title", "Similarity"])
    shown = dups.get("pairs", [])[:MAX_DUP_ROWS]
    for pr in shown:
        newer, older = max(pr["a"], pr["b"]), min(pr["a"], pr["b"])
        ws_dup.append([newer, titles.get(newer), older, titles.get(older), pr["similarity"]])
    url_base = f"https://{host}/{repo}/issues/"
    for cells in ws_dup.iter_rows(min_row=2, max_row=len(shown) + 1, max_col=5):
        for c in cells:
            c.font = base
        for c in (cells[0], cells[2]):
            c.hyperlink, c.font = url_base + str(c.value), link
    for col, width in zip("ABCDE", (11, 60, 11, 60, 11)):
        ws_dup.column_dimensions[col].width = width
    ws_dup.freeze_panes = "A2"
    ws_dup.cell(row=len(shown) + 3, column=1,
                value=f"Method: {dups.get('method') or 'not run'}; threshold {dups.get('threshold')}; showing "
                      f"{len(shown)} of {len(dups.get('pairs', []))} pairs. Candidates only: read both issues "
                      "before closing one.").font = Font(name="Arial", italic=True, size=9)

    # Labels: how each label was interpreted, so the profile can be tuned
    ws_lab = wb.create_sheet("Labels")
    header(ws_lab, ["Label", "Open issues", "How it's used"])
    label_counts = Counter(l for i in blob["issues"] for l in i["labels"])
    for label, n in sorted(label_counts.items(), key=lambda kv: (-kv[1], kv[0])):
        ws_lab.append([label, n, profile.describe_label(label) or "(not used)"])
    for cells in ws_lab.iter_rows(min_row=2, max_row=len(label_counts) + 1, max_col=3):
        for c in cells:
            c.font = base
    for col, width in zip("ABC", (40, 12, 60)):
        ws_lab.column_dimensions[col].width = width
    ws_lab.freeze_panes = "A2"

    # Scoring explanation
    ws_sc = wb.create_sheet("Scoring")
    lines = [("How issues are scored", None), ("", None),
             ("Score = priority points + severity points + signal points. Higher = fix sooner.", None), ("", None),
             ("Priority", "Points")]
    lines += [(k if k != "none" else "No priority label", v) for k, v in w["priority"].items()]
    lines += [("", None), ("Severity (bugs only; worst of labels and AI)", "Points")]
    lines += [(k, v) for k, v in w["severity"].items()]
    lines += [("", None), ("Signals", "Points")]
    lines += [(k.replace("_", " "), v) for k, v in w["signals"].items()]
    lines += [("thumbs-up reactions", f"{w['thumbs_up_scale']:g} × log2(1 + count), max {w['thumbs_up_cap']:g}"),
              ("", None), ("Urgency tier", "Minimum score")]
    lines += [(k, v) for k, v in w["urgency_thresholds"].items()] + [("Low", "anything lower")]
    lines += [(f"{k} floor", f"at least {v}") for k, v in w["urgency_floors"].items()]
    lines += [("", None), ("Rules", None),
              ("Human labels win", "Labels decide the kind; the AI only fills in issues without a kind label, and "
                                   "flags disagreements it is confident about."),
              ("AI fills a gap", f"when its confidence is at least {w['ai_min_confidence_fill']}"),
              ("AI flags a disagreement", f"when its confidence is at least {w['ai_min_confidence_flag']}"),
              ("Unlabeled issues", "treated as bugs (profile setting)" if profile.uncategorized_is_bug
               else "left as Uncategorized unless the AI classifies them"),
              ("No labels and no AI", "a title mentioning a crash, hang, freeze, segfault, ANR or data loss "
                                      "is rated high severity"),
              ("Label profile", profile.source),
              ("Recent releases", ", ".join(f"{a}.{b}" for a, b in sorted(ctx.recent_releases, reverse=True)) or "none"),
              ("", None), ("Data", None), ("Repository", f"{host}/{repo}"),
              ("Exported at (UTC)", blob.get("exported_at")), ("AI-classified issues", len(ai)),
              ("Duplicate method", dups.get("method") or "not run")]
    for a, b in lines:
        ws_sc.append([a, b])
    for r in ws_sc.iter_rows():
        for c in r:
            c.font = base
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if r[1].value in ("Points", "Minimum score") or r[0].value in ("Rules", "Data"):
            r[0].font = r[1].font = bold
    ws_sc["A1"].font = title_font
    ws_sc.column_dimensions["A"].width = 46
    ws_sc.column_dimensions["B"].width = 80

    # Summary: live formulas over All Issues
    n = len(all_rows) + 1

    def rng(key):
        L = get_column_letter(COL_INDEX[key] + 1)
        return f"'All Issues'!${L}$2:${L}${max(n, 2)}"

    kind_r, urg_r, team_r, stat_r, pri_r, sev_r, flag_r = (rng(k) for k in (
        "kind", "urgency", "team", "triage_status", "priority", "severity", "flags"))
    ws_sum["A1"] = f"Open issue triage: {repo}"
    ws_sum["A1"].font = title_font
    ws_sum["A2"] = (f"Exported {blob.get('exported_at')} UTC. {len(ai):,} issues AI-classified. "
                    f"Label profile: {profile.source}. Duplicates: {dups.get('method') or 'not run'}.")
    ws_sum["A2"].font = Font(name="Arial", size=9, italic=True)
    ws_sum["A4"], ws_sum["B4"] = "Total open issues", f"=COUNTA('All Issues'!$A$2:$A${max(n, 2)})"

    def section(row, title, key_range, values, patterns=None):
        ws_sum.cell(row=row, column=1, value=title).font = bold
        ws_sum.cell(row=row, column=2, value="Issues").font = bold
        for k, v in enumerate(values, 1):
            ws_sum.cell(row=row + k, column=1, value=v)
            if patterns:
                f = '=COUNTIF(' + key_range + ',"*' + patterns[k - 1] + '*")'
            else:
                crit = '""' if v == "(none)" else '"' + v + '"'
                f = "=COUNTIFS(" + key_range + "," + crit + ")"
            ws_sum.cell(row=row + k, column=2, value=f)
        return row + len(values) + 2

    row = section(6, "Kind", kind_r, list(KIND_DISPLAY.values()))
    row = section(row, "Triage status", stat_r, ["Triaged", "Needs triage", "Needs primary triage",
                                                  "Needs team triage", "Triaged, no priority", "Waiting for reporter"])
    row = section(row, "Priority", pri_r, ["P0", "P1", "P2", "P3", "(none)"])
    row = section(row, "Bug severity", sev_r, ["Critical", "High", "Medium", "Low", "Unrated"])
    row = section(row, "Review flags", flag_r, [a for a, _ in FLAG_SUMMARY], [b for _, b in FLAG_SUMMARY])

    ws_sum.cell(row=row, column=1, value="Bugs by team / area and urgency").font = bold
    for c, t in enumerate(list(URGENCY_TIERS) + ["Total"], 2):
        ws_sum.cell(row=row, column=c, value=t).font = bold
    team_counts = Counter(r["team"] for r in bugs if r["team"])
    teams = [t for t, _ in team_counts.most_common(MAX_TEAMS_IN_SUMMARY)] + ["(none)"]
    for k, team in enumerate(teams, 1):
        rr = row + k
        ws_sum.cell(row=rr, column=1, value=team)
        crit = '""' if team == "(none)" else '"' + team.replace('"', '""') + '"'
        for c, t in enumerate(URGENCY_TIERS, 2):
            ws_sum.cell(row=rr, column=c, value=f'=COUNTIFS({team_r},{crit},{urg_r},"{t}")')
        ws_sum.cell(row=rr, column=len(URGENCY_TIERS) + 2, value=f"=SUM(B{rr}:{get_column_letter(len(URGENCY_TIERS) + 1)}{rr})")
    total_row = row + len(teams) + 1
    ws_sum.cell(row=total_row, column=1, value="Total" if len(team_counts) <= MAX_TEAMS_IN_SUMMARY
                else f"Total (top {MAX_TEAMS_IN_SUMMARY} teams shown)").font = bold
    for c in range(2, len(URGENCY_TIERS) + 3):
        L = get_column_letter(c)
        ws_sum.cell(row=total_row, column=c, value=f"=SUM({L}{row + 1}:{L}{total_row - 1})").font = bold
    for r in ws_sum.iter_rows(min_row=3):
        for c in r:
            if not c.font.b:
                c.font = base
    ws_sum.column_dimensions["A"].width = 38
    for L in "BCDEF":
        ws_sum.column_dimensions[L].width = 12

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    with out_path.with_suffix(".csv").open("w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f)
        wr.writerow([h for h, _, _ in COLUMNS] + ["URL"])
        for r in all_rows:
            wr.writerow([r.get(k) for _, k, _ in COLUMNS] + [r["url"]])

    return {
        "issues": len(rows), "bugs": len(bugs), "review": len(review), "dup_pairs": len(dups.get("pairs", [])),
        "kinds": dict(Counter(r["kind"] for r in rows)),
        "urgency": {t: sum(1 for r in bugs if r["urgency"] == t) for t in URGENCY_TIERS},
        "uncategorized": sum(1 for r in rows if r["kind"] == "Uncategorized"),
        "xlsx": out_path, "csv": out_path.with_suffix(".csv"),
    }
