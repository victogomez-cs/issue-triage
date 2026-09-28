"""End to end through the command line, with GitHub and Anthropic faked."""

import json

import pytest
import requests
from openpyxl import load_workbook

from issue_triage import cli


@pytest.fixture
def env(monkeypatch, tmp_path, sample_issues, fake_github, fake_anthropic):
    gh = fake_github(sample_issues)
    monkeypatch.setattr(requests.Session, "post", lambda self, *a, **k: gh.post(*a, **k))
    monkeypatch.setenv("GITHUB_TOKEN", "t")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    import anthropic
    monkeypatch.setattr(anthropic, "Anthropic", fake_anthropic)
    monkeypatch.chdir(tmp_path)
    return tmp_path


def test_run_without_ai(env, capsys):
    cli.main(["https://github.com/o/r/issues"])          # shortcut for `run`
    report = env / "triage-output" / "o__r" / "o__r_triage.xlsx"
    assert report.exists() and report.with_suffix(".csv").exists()
    wb = load_workbook(report)
    assert wb.sheetnames == ["Summary", "Bugs Ranked", "Review Queue", "All Issues", "Duplicates", "Labels", "Scoring"]
    bugs = [r for r in wb["Bugs Ranked"].iter_rows(min_row=2, values_only=True)]
    assert bugs[0][1] == 1  # the crash with priority: high is ranked first
    labels = {r[0]: r[2] for r in wb["Labels"].iter_rows(min_row=2, values_only=True)}
    assert labels["priority: high"] == "priority: P1" and "team/area: ui" in labels["area/ui"]
    err = capsys.readouterr().err
    assert "no kind label" in err and "AI classification off" in err


def test_run_with_ai(env, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "k")
    cli.main(["run", "o/r", "--yes"])
    data = env / "triage-output" / "o__r"
    assert len([l for l in (data / "ai_results.jsonl").read_text().splitlines() if l]) == 10
    wb = load_workbook(data / "o__r_triage.xlsx")
    rows = {r[0]: r for r in wb["All Issues"].iter_rows(min_row=2, values_only=True)}
    assert rows[6][2] == "Feature / proposal" or rows[6][2] == "Bug"   # AI filled the unlabeled issue
    # second run reuses the export and the AI cache
    cli.main(["run", "o/r", "--yes"])


def test_import_then_report(env, tmp_path):
    cli.main(["export", "o/r"])
    f = tmp_path / "page.json"
    f.write_text(json.dumps([{"number": 6, "kind": "bug", "severity": "critical", "confidence": 0.9,
                              "summary": "From the page."}]))
    cli.main(["import-ai", "o/r", str(f)])
    cli.main(["report", "o/r"])
    wb = load_workbook(env / "triage-output" / "o__r" / "o__r_triage.xlsx")
    rows = {r[0]: r for r in wb["All Issues"].iter_rows(min_row=2, values_only=True)}
    assert rows[6][3] == "Critical" and rows[6][18] == "From the page."


def test_bad_repo_exits_cleanly(env):
    with pytest.raises(SystemExit) as e:
        cli.main(["run", "not a repo"])
    assert e.value.code == 1


def test_labels_and_profiles_commands(env, capsys):
    cli.main(["export", "o/r"])
    cli.main(["labels", "o/r"])
    assert "priority: P1" in capsys.readouterr().out
    cli.main(["profiles"])
    assert "flutter" in capsys.readouterr().out
